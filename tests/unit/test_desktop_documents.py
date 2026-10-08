from __future__ import annotations

from pathlib import Path
from typing import cast

import pytest

pytest.importorskip("PySide6")

from PySide6.QtWidgets import (
    QApplication,
    QComboBox,
    QLineEdit,
    QMessageBox,
    QPushButton,
    QTableWidget,
    QTabWidget,
)

from school_timetable_solver.adapter.project_store_adapter import LocalProjectStoreAdapter
from school_timetable_solver.desktop_composition import DesktopApplicationComposition
from school_timetable_solver.model.input_models import InputDataModel
from school_timetable_solver.service.project_document_services import ImportProjectDocumentService
from school_timetable_solver.ui.editor_navigation import EditorSection


@pytest.fixture
def desktop_app() -> QApplication:
    existing = QApplication.instance()
    return cast(QApplication, existing) if existing is not None else QApplication([])


def test_editor_saves_count_explicitly_and_uses_ids_for_renamed_duplicate_teachers(
    desktop_app: QApplication, tmp_path: Path, minimal_input_data: InputDataModel
) -> None:
    window = DesktopApplicationComposition().create_desktop_window(tmp_path)
    window._create_new_project()
    project_id = window._current_project_id
    assert project_id is not None
    editor = window._seasonal_editor()
    revision = editor.document.revision
    editor.document = ImportProjectDocumentService().execute(minimal_input_data)
    editor.document.revision = revision
    editor.document_changed.emit()
    assert window._save_current()
    editor._select_section(EditorSection.LESSON_COUNTS)
    table = editor.findChild(QTableWidget, "lesson_counts")
    assert table is not None
    cast(QLineEdit, table.cellWidget(0, 1)).setText("3")
    assert table.item(0, table.columnCount() - 1).text() == "3"
    stored = LocalProjectStoreAdapter(tmp_path).load_document(project_id)
    assert stored is not None and stored.lesson_requirements[0].required_periods != "3"
    save = editor.findChild(QPushButton, "page_save")
    assert save is not None
    save.click()
    stored = LocalProjectStoreAdapter(tmp_path).load_document(project_id)
    assert stored is not None and stored.lesson_requirements[0].required_periods == "3"
    editor.document.teachers[1].teacher_name = editor.document.teachers[0].teacher_name
    editor.document_changed.emit()
    assert window._save_current()
    editor._select_section(EditorSection.TEACHER_ASSIGNMENTS)
    table = editor.findChild(QTableWidget, "teacher_assignments")
    assert table is not None
    combo = cast(QComboBox, table.cellWidget(0, 2))
    assert combo.currentData() == "T1"
    assert combo.itemData(1) != combo.itemData(2)
    combo.setCurrentIndex(combo.findData("T2"))
    assert window._save_current()
    window._show_home()
    window._open_project(project_id)
    assert window._seasonal_editor().document.lesson_requirements[0].teacher_id == "T2"
    window.close()


def test_master_creation_and_removal_preserve_references_and_incomplete_drafts(
    desktop_app: QApplication, tmp_path: Path
) -> None:
    window = DesktopApplicationComposition().create_desktop_window(tmp_path)
    window._create_new_project()
    editor = window._seasonal_editor()
    editor._select_section(EditorSection.MASTER)
    editor._append_row("campuses")
    campus_table = editor.findChild(QTableWidget, "campuses")
    assert campus_table is not None
    cast(QLineEdit, campus_table.cellWidget(0, 0)).setText("新校舎")
    campus_id = editor.document.campuses[0].campus_id
    editor._append_row("rooms")
    rooms_table = editor.findChild(QTableWidget, "rooms")
    assert rooms_table is not None
    campus_combo = cast(QComboBox, rooms_table.cellWidget(0, 1))
    campus_combo.setCurrentIndex(campus_combo.findData(campus_id))
    editor._remove_row("campuses", 0)
    assert editor.document.rooms[0].campus_id == campus_id
    assert editor.document.rooms[0].room_name == ""
    assert editor.has_unsaved_changes
    assert window._save_current()
    assert editor.document.revision > 1
    assert not editor.has_unsaved_changes
    window.close()


def test_schedule_creation_preserves_existing_day_edits_and_period_ids(
    desktop_app: QApplication, tmp_path: Path, minimal_input_data: InputDataModel
) -> None:
    window = DesktopApplicationComposition().create_desktop_window(tmp_path)
    window._create_new_project()
    editor = window._seasonal_editor()
    revision = editor.document.revision
    editor.document = ImportProjectDocumentService().execute(minimal_input_data)
    editor.document.revision = revision
    editor.document_changed.emit()
    assert window._save_current()
    editor._set_day_periods(0, False)
    editor.document.end_date = "2026-07-31"
    editor._create_dates()
    assert editor.document.calendar_days[0].enabled_period_ids == ""
    assert len(editor.document.calendar_days) == 5
    assert editor.document.calendar_days[-1].target_date == "2026-07-31"
    assert editor.document.calendar_days[-1].enabled_period_ids == "P1|P2|P3|P4|P5|P6"
    assert window._save_current()
    window.close()


def test_save_conflict_keeps_editor_open_and_dirty(
    desktop_app: QApplication, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    window = DesktopApplicationComposition().create_desktop_window(tmp_path)
    window._create_new_project()
    editor = window._seasonal_editor()
    project_id = window._current_project_id
    assert project_id is not None
    store = LocalProjectStoreAdapter(tmp_path)
    other = store.load_document(project_id)
    assert other is not None
    store.save_document(project_id, other)
    editor.document.start_date = "2027-01-01"
    editor.document_changed.emit()
    assert not window._save_current()
    assert editor.has_unsaved_changes
    assert "未保存" in editor._save_status.text()
    monkeypatch.setattr(QMessageBox, "exec", lambda *args: QMessageBox.StandardButton.No)
    window._show_home()
    assert window._stack.currentWidget() is editor
    editor.discard_changes()
    window.close()


def test_blank_project_widgets_can_produce_a_verified_timetable(
    desktop_app: QApplication, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from openpyxl import load_workbook
    from PySide6.QtCore import QDate, QEventLoop, QTime, QTimer
    from PySide6.QtWidgets import QDateEdit, QTimeEdit

    from school_timetable_solver.model.input_models import GenerationMode
    from school_timetable_solver.model.project_models import ProjectExecutionSettingsModel

    window = DesktopApplicationComposition().create_desktop_window(tmp_path)
    window._create_new_project()
    editor = window._seasonal_editor()
    editor._select_section(EditorSection.MASTER)
    for group in ("campuses", "rooms", "teachers", "classes", "subjects"):
        editor._append_row(group)
        table = editor.findChild(QTableWidget, group)
        assert table is not None
        cast(QLineEdit, table.cellWidget(0, 0)).setText(group)
        if group in ("rooms", "teachers", "classes"):
            cast(QComboBox, table.cellWidget(0, 1)).setCurrentIndex(1)
        if group == "classes":
            cast(QLineEdit, table.cellWidget(0, 3)).setText("1")
    for index in range(6):
        editor._append_row("periods")
        table = editor.findChild(QTableWidget, "periods")
        assert table is not None
        cast(QLineEdit, table.cellWidget(index, 0)).setText(str(index + 1))
        cast(QTimeEdit, table.cellWidget(index, 1)).setTime(QTime(index + 8, 0))
        cast(QTimeEdit, table.cellWidget(index, 2)).setTime(QTime(index + 9, 0))
    assert window._save_current()
    editor._select_section(EditorSection.SCHEDULE)
    first = editor.findChild(QDateEdit, "start_date")
    last = editor.findChild(QDateEdit, "end_date")
    assert first is not None and last is not None
    first.setDate(QDate(2027, 1, 1))
    last.setDate(QDate(2027, 1, 1))
    editor._create_dates()
    assert window._save_current()
    editor._select_section(EditorSection.LESSON_COUNTS)
    counts = editor.findChild(QTableWidget, "lesson_counts")
    assert counts is not None
    cast(QLineEdit, counts.cellWidget(0, 1)).setText("1")
    assert window._save_current()
    editor._select_section(EditorSection.TEACHER_ASSIGNMENTS)
    assignments = editor.findChild(QTableWidget, "teacher_assignments")
    assert assignments is not None
    cast(QComboBox, assignments.cellWidget(0, 2)).setCurrentIndex(1)
    # A class policy is a required user input, not an implicit solver default.
    assert window._save_current()
    editor._select_section(EditorSection.PLACEMENT_CONDITIONS)
    editor._append_row("placement_rules")
    rules = editor.findChild(QTableWidget, "placement_rules")
    assert rules is not None
    cast(QLineEdit, rules.cellWidget(0, 0)).setText("クラス基本方針")
    cast(QComboBox, rules.cellWidget(0, 2)).setCurrentIndex(1)
    cast(QComboBox, rules.cellWidget(0, 3)).setCurrentIndex(1)
    # Pick the first period through the same model binding as the checklist widget.
    row = editor.document.placement_rules[0]
    editor._set_value(row, "allowed_period_ids", editor.document.periods[0].period_id)
    editor._set_value(row, "priority", "10")
    editor._set_value(row, "daily_hard_limit", "1")
    editor._append_row("placement_rules")
    teacher_policy = editor.document.placement_rules[1]
    editor._set_value(teacher_policy, "rule_name", "教師基本方針")
    editor._set_value(teacher_policy, "constraint_type", "hard")
    editor._set_value(teacher_policy, "target_entity", "teacher")
    editor._set_value(teacher_policy, "daily_hard_limit", "1")
    editor._set_value(teacher_policy, "forbid_first_last_same_day", "false")
    editor._set_value(teacher_policy, "priority", "10")
    assert window._save_current()
    output = tmp_path / "new-project-result.xlsx"
    monkeypatch.setattr(QMessageBox, "information", lambda *args: QMessageBox.StandardButton.Ok)
    monkeypatch.setattr(QMessageBox, "critical", lambda *args: QMessageBox.StandardButton.Ok)
    monkeypatch.setattr(QMessageBox, "warning", lambda *args: QMessageBox.StandardButton.Ok)
    project_id = window._current_project_id
    assert project_id is not None
    window._start_project_run(
        project_id, ProjectExecutionSettingsModel(output, None, GenerationMode.STRICT, 5, 1, 1)
    )
    loop = QEventLoop()
    assert window._run_thread is not None
    window._run_thread.finished.connect(loop.quit)
    QTimer.singleShot(10000, loop.quit)
    loop.exec()
    desktop_app.processEvents()
    assert window._run_thread is None
    assert output.is_file(), [(issue.rule_id, issue.message) for issue in editor._issues]
    workbook = load_workbook(output, data_only=True)
    assert workbook.sheetnames == ["全体"]
    assert editor._issues == ()
    window.close()


@pytest.mark.parametrize("action", ["page", "home", "close", "reload"])
@pytest.mark.parametrize("discard", [False, True])
def test_unsaved_confirmation_preserves_or_discards_edits(
    desktop_app: QApplication,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    action: str,
    discard: bool,
) -> None:
    from PySide6.QtCore import QDate
    from PySide6.QtWidgets import QDateEdit

    window = DesktopApplicationComposition().create_desktop_window(tmp_path)
    window._create_new_project()
    window.show()
    editor = window._seasonal_editor()
    project_id = window._current_project_id
    assert project_id is not None
    store = LocalProjectStoreAdapter(tmp_path)
    baseline = store.load_document(project_id)
    assert baseline is not None
    first = editor.findChild(QDateEdit, "start_date")
    assert first is not None
    first.setDate(QDate(2027, 1, 1))
    assert editor.has_unsaved_changes
    answers: list[str] = []

    def answer(dialog: QMessageBox) -> int:
        answers.append(dialog.objectName())
        assert dialog.button(QMessageBox.StandardButton.Yes).text() == "はい"
        assert dialog.button(QMessageBox.StandardButton.No).text() == "いいえ"
        assert dialog.defaultButton() == dialog.button(QMessageBox.StandardButton.No)
        return int(QMessageBox.StandardButton.Yes if discard else QMessageBox.StandardButton.No)

    monkeypatch.setattr(QMessageBox, "exec", answer)
    if action == "page":
        editor._navigation_buttons[EditorSection.LESSON_COUNTS].click()
        expected = EditorSection.LESSON_COUNTS if discard else EditorSection.SCHEDULE
        assert editor._section == expected
        assert editor._navigation_buttons[expected].isChecked()
    elif action == "home":
        window._show_home()
        assert window._stack.currentWidget() == (window._home if discard else editor)
        assert window._current_project_id == (None if discard else project_id)
    elif action == "close":
        assert window.close() == discard
        assert window.isVisible() != discard
    else:
        window._reload_current()
        assert editor._section == EditorSection.SCHEDULE
    assert answers == ["unsaved_changes_dialog"]
    assert store.load_document(project_id) == baseline
    assert editor.has_unsaved_changes != discard
    assert editor.document.start_date == (baseline.start_date if discard else "2027-01-01")
    if not discard:
        assert editor.findChild(QDateEdit, "start_date") is first
        assert first.date() == QDate(2027, 1, 1)
    editor.discard_changes()
    window.close()


@pytest.mark.parametrize("section", [EditorSection.MASTER, EditorSection.PLACEMENT_CONDITIONS])
@pytest.mark.parametrize("discard", [False, True])
def test_unsaved_tab_change_is_guarded(
    desktop_app: QApplication,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    section: EditorSection,
    discard: bool,
) -> None:
    window = DesktopApplicationComposition().create_desktop_window(tmp_path)
    window._create_new_project()
    editor = window._seasonal_editor()
    editor._select_section(section)
    baseline = editor.document.revision
    editor._append_row("campuses" if section == EditorSection.MASTER else "placement_rules")
    tabs = editor.findChild(QTabWidget)
    assert tabs is not None
    monkeypatch.setattr(
        QMessageBox,
        "exec",
        lambda *args: QMessageBox.StandardButton.Yes if discard else QMessageBox.StandardButton.No,
    )
    tabs.setCurrentIndex(1)
    current_tabs = editor.findChild(QTabWidget)
    assert current_tabs is not None
    assert current_tabs.currentIndex() == (1 if discard else 0)
    assert editor.document.revision == baseline
    assert editor.has_unsaved_changes != discard
    assert len(
        editor.document.campuses
        if section == EditorSection.MASTER
        else editor.document.placement_rules
    ) == (0 if discard else 1)
    editor.discard_changes()
    window.close()


def test_cancel_restores_last_save_and_execution_does_not_save_implicitly(
    desktop_app: QApplication,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    minimal_input_data: InputDataModel,
) -> None:
    window = DesktopApplicationComposition().create_desktop_window(tmp_path)
    window._create_new_project()
    editor = window._seasonal_editor()
    revision = editor.document.revision
    editor.document = ImportProjectDocumentService().execute(minimal_input_data)
    editor.document.revision = revision
    editor.document_changed.emit()
    assert window._save_current()
    editor._select_section(EditorSection.LESSON_COUNTS)
    table = editor.findChild(QTableWidget, "lesson_counts")
    assert table is not None
    cast(QLineEdit, table.cellWidget(0, 1)).setText("3")
    editor._page_save.click()
    project_id = window._current_project_id
    assert project_id is not None
    store = LocalProjectStoreAdapter(tmp_path)
    baseline = store.load_document(project_id)
    assert baseline is not None and baseline.lesson_requirements[0].required_periods == "3"
    cast(QLineEdit, table.cellWidget(0, 1)).setText("4")
    warnings: list[str] = []
    monkeypatch.setattr(
        QMessageBox, "warning", lambda _parent, _title, message: warnings.append(message)
    )
    window._validate_current()
    window._run_current()
    assert len(warnings) == 2
    assert window._run_thread is None
    assert store.load_document(project_id) == baseline
    editor._page_cancel.click()
    assert editor.document == baseline
    assert not editor.has_unsaved_changes
    table = editor.findChild(QTableWidget, "lesson_counts")
    assert table is not None
    assert cast(QLineEdit, table.cellWidget(0, 1)).text() == "3"
    assert table.item(0, table.columnCount() - 1).text() == "3"
    assert not editor._page_save.isEnabled()
    assert not editor._page_cancel.isEnabled()
    window._show_home()
    window._open_project(project_id)
    assert editor.document == baseline
    window.close()


def test_closing_confirmation_keeps_unsaved_page(
    desktop_app: QApplication,
    tmp_path: Path,
) -> None:
    from PySide6.QtCore import QDate, Qt, QTimer
    from PySide6.QtTest import QTest
    from PySide6.QtWidgets import QDateEdit

    window = DesktopApplicationComposition().create_desktop_window(tmp_path)
    window._create_new_project()
    editor = window._seasonal_editor()
    first = editor.findChild(QDateEdit, "start_date")
    assert first is not None
    first.setDate(QDate(2027, 1, 1))
    dialogs: list[str] = []

    def dismiss() -> None:
        dialog = desktop_app.activeModalWidget()
        if isinstance(dialog, QMessageBox):
            dialogs.append(dialog.objectName())
            QTest.keyClick(dialog, Qt.Key.Key_Escape)

    timed_out: list[bool] = []

    def stop_dialog() -> None:
        timed_out.append(True)
        dialog = desktop_app.activeModalWidget()
        if isinstance(dialog, QMessageBox):
            dialog.reject()

    timeout = QTimer()
    timeout.setSingleShot(True)
    timeout.timeout.connect(stop_dialog)
    timeout.start(1000)
    QTimer.singleShot(0, dismiss)
    editor._select_section(EditorSection.LESSON_COUNTS)
    timeout.stop()
    assert not timed_out
    assert dialogs == ["unsaved_changes_dialog"]
    assert editor._section == EditorSection.SCHEDULE
    assert editor.has_unsaved_changes
    assert first.date() == QDate(2027, 1, 1)
    editor.discard_changes()
    window.close()
